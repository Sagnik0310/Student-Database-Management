import csv
import io
import os
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from database.db_connection import get_connection
from logs.logger import logger
from services.academic_service import (
    calculate_student_gpa,
    can_manage_course,
    compute_attendance_metrics,
    get_course,
    get_grading_scale,
    list_attendance,
    list_courses,
    list_enrollments,
    list_grades,
)
from services.auth_service import can_access_student, filter_students_for_user
from services.student_service import search_student, view_students

DATE_REGEX = re.compile(r"^\d{4}-\d{2}-\d{2}$")
FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@", "\t", "\r", "\n")

VALID_REPORT_TYPES = {
    "student_directory": "Student Directory & Academic Standing Report",
    "course_enrollment": "Course Enrollment & Capacity Report",
    "attendance_register": "Attendance Register & Compliance Report",
    "grade_sheet": "Course Assessment & Grade Sheet Report",
    "academic_transcript": "Student Academic Transcript & Credit-Weighted CGPA",
    "executive_summary": "Institutional Executive Summary (Branch, Semester & Course)",
}

VALID_EXPORT_FORMATS = {"csv", "xlsx", "excel", "pdf"}


def _get_max_export_rows():
    try:
        return max(10, min(10000, int(os.environ.get("MAX_EXPORT_ROWS", 2000))))
    except (ValueError, TypeError):
        return 2000


def _get_max_date_range_days():
    try:
        return max(1, min(3650, int(os.environ.get("MAX_DATE_RANGE_DAYS", 366))))
    except (ValueError, TypeError):
        return 366


def _is_test_mode():
    return os.environ.get("STUDENTDB_TEST_MODE") == "1"


def _get_db():
    if _is_test_mode():
        return None
    return get_connection()


# ==========================================
# SECURITY HELPERS: FORMULA INJECTION & SAFE FILENAMES
# ==========================================
def sanitize_spreadsheet_cell(value):
    """
    Mitigate CSV/Excel formula injection (CSV/DDE injection).
    Any string starting with '=', '+', '-', '@', tab, or carriage return/newline
    (including after leading spaces) is prefixed with a single quote (').
    Numeric values (int, float) are preserved as numbers.
    """
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value

    text = str(value)
    if not text:
        return ""

    stripped = text.lstrip(" ")
    if text.startswith(FORMULA_TRIGGER_CHARS) or (
        stripped and stripped.startswith(FORMULA_TRIGGER_CHARS)
    ):
        return f"'{text}"
    return text


def build_safe_filename(report_type, export_format, suffix=""):
    """
    Construct a traversal-safe filename for Content-Disposition headers.
    Strips path separators, '..' sequences, and non-alphanumeric characters.
    """
    clean_type = re.sub(r"[^A-Za-z0-9_-]", "_", str(report_type or "report"))
    clean_type = re.sub(r"_+", "_", clean_type).strip("_") or "report"

    clean_suffix = ""
    if suffix:
        clean_suffix = "_" + re.sub(r"[^A-Za-z0-9_-]", "_", str(suffix))
        clean_suffix = re.sub(r"_+", "_", clean_suffix).rstrip("_")

    fmt = str(export_format or "csv").lower().strip()
    ext = "xlsx" if fmt in ("xlsx", "excel") else ("pdf" if fmt == "pdf" else "csv")
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    return f"studentdb_{clean_type}{clean_suffix}_{ts}.{ext}"


# ==========================================
# FILTER & RBAC VALIDATION
# ==========================================
def validate_and_normalize_filters(raw_filters=None):
    """
    Validate report/analytics query parameters server-side:
    - Validates date_from / date_to format (YYYY-MM-DD), ordering, and max span.
    - Caps limit to MAX_EXPORT_ROWS.
    - Sanitizes string filters.
    """
    params = dict(raw_filters or {})
    max_rows = _get_max_export_rows()
    max_days = _get_max_date_range_days()

    report_type = str(params.get("report_type") or "student_directory").strip().lower()
    if report_type not in VALID_REPORT_TYPES:
        return {
            "error": f"Invalid report_type '{report_type}'. Must be one of: {', '.join(sorted(VALID_REPORT_TYPES))}"
        }

    export_format = str(params.get("format") or "csv").strip().lower()
    if export_format not in VALID_EXPORT_FORMATS:
        return {
            "error": f"Invalid export format '{export_format}'. Must be one of: csv, xlsx, pdf"
        }

    branch = str(params.get("branch") or "").strip()
    if branch.lower() == "all":
        branch = ""
    elif branch:
        branch = branch.upper()

    semester = str(params.get("semester") or "").strip()
    if semester.lower() == "all":
        semester = ""

    academic_year = str(params.get("academic_year") or "").strip()
    if academic_year.lower() == "all":
        academic_year = ""

    course_code = str(params.get("course_code") or "").strip().upper()
    if course_code.lower() == "all":
        course_code = ""

    reg_no = str(params.get("reg_no") or "").strip()
    if reg_no.lower() == "all":
        reg_no = ""

    status = str(params.get("status") or "").strip()
    if status.lower() == "all":
        status = ""

    date_from = str(params.get("date_from") or "").strip()
    date_to = str(params.get("date_to") or "").strip()

    dt_from_obj = None
    dt_to_obj = None
    if date_from:
        if not DATE_REGEX.match(date_from):
            return {"error": "date_from must be in YYYY-MM-DD format"}
        try:
            dt_from_obj = datetime.strptime(date_from, "%Y-%m-%d")
        except ValueError:
            return {"error": "date_from is not a valid calendar date"}

    if date_to:
        if not DATE_REGEX.match(date_to):
            return {"error": "date_to must be in YYYY-MM-DD format"}
        try:
            dt_to_obj = datetime.strptime(date_to, "%Y-%m-%d")
        except ValueError:
            return {"error": "date_to is not a valid calendar date"}

    if dt_from_obj and dt_to_obj:
        if dt_from_obj > dt_to_obj:
            return {"error": "date_from cannot be later than date_to"}
        span_days = (dt_to_obj - dt_from_obj).days
        if span_days > max_days:
            return {
                "error": f"Date range ({span_days} days) exceeds maximum allowed span of {max_days} days"
            }

    limit_raw = params.get("limit")
    if limit_raw is not None and str(limit_raw).strip() != "":
        try:
            limit_val = int(limit_raw)
        except (ValueError, TypeError):
            return {"error": "limit must be a valid integer"}
        if limit_val < 1:
            return {"error": "limit must be at least 1"}
        if limit_val > max_rows:
            return {
                "error": f"Requested limit ({limit_val}) exceeds maximum export limit of {max_rows} rows"
            }
    else:
        limit_val = min(500, max_rows)

    return {
        "report_type": report_type,
        "format": "xlsx" if export_format == "excel" else export_format,
        "branch": branch,
        "semester": semester,
        "academic_year": academic_year,
        "course_code": course_code,
        "reg_no": reg_no,
        "status": status,
        "date_from": date_from,
        "date_to": date_to,
        "limit": limit_val,
    }


def enforce_report_rbac(user, normalized_filters):
    """
    Validate that the current user is authorized for the requested report scope
    and return a copy of normalized_filters with role boundaries applied.
    """
    if not user or not isinstance(user, dict):
        return {"error": "Authentication required", "status_code": 401}

    role = user.get("role")
    scoped = dict(normalized_filters)

    if role == "Student":
        linked_reg = str(user.get("linked_reg_no") or "").strip()
        if not linked_reg:
            return {
                "error": "Access denied: Student account has no linked registration number",
                "status_code": 403,
            }
        if scoped.get("report_type") == "executive_summary":
            return {
                "error": "Access denied: Institutional executive summaries are restricted to Admin and Faculty",
                "status_code": 403,
            }
        if scoped.get("reg_no") and scoped["reg_no"] != linked_reg:
            return {
                "error": "Access denied: Students may only generate reports or analytics for their own registration number",
                "status_code": 403,
            }
        scoped["reg_no"] = linked_reg
        return {"filters": scoped}

    if role == "Faculty":
        assigned_branches = [
            str(b).strip().upper() for b in (user.get("assigned_branches") or [])
        ]
        wildcard = "*" in assigned_branches

        if scoped.get("branch") and not wildcard:
            if scoped["branch"].upper() not in assigned_branches:
                return {
                    "error": f"Access denied: Faculty is not assigned to branch '{scoped['branch']}'",
                    "status_code": 403,
                }

        if scoped.get("course_code"):
            course = get_course(scoped["course_code"])
            if not course:
                return {
                    "error": f"Course '{scoped['course_code']}' not found",
                    "status_code": 404,
                }
            if not can_manage_course(user, course):
                return {
                    "error": f"Access denied: You are not assigned to course '{scoped['course_code']}'",
                    "status_code": 403,
                }

        if scoped.get("reg_no"):
            student = search_student(scoped["reg_no"])
            if not student:
                return {
                    "error": f"Student '{scoped['reg_no']}' not found",
                    "status_code": 404,
                }
            if not can_access_student(user, student):
                return {
                    "error": f"Access denied: You are not assigned to student '{scoped['reg_no']}'",
                    "status_code": 403,
                }

        return {"filters": scoped}

    if role == "Admin":
        return {"filters": scoped}

    return {"error": "Access denied: Unrecognized role", "status_code": 403}


# ==========================================
# OBJECTIVE A: ADVANCED ANALYTICS WORKSPACE
# ==========================================
def _run_mongo_branch_pipeline(match_query):
    """Run bounded MongoDB aggregation pipeline for branch statistics when MongoDB is active."""
    db = _get_db()
    if db is None:
        return None
    try:
        pipeline = [
            {"$match": match_query},
            {
                "$group": {
                    "_id": {"$ifNull": ["$branch", "Unassigned"]},
                    "student_count": {"$sum": 1},
                    "avg_age": {"$avg": "$age"},
                    "avg_gpa": {"$avg": "$gpa"},
                    "avg_attendance": {"$avg": "$attendance_pct"},
                }
            },
            {"$sort": {"student_count": -1, "_id": 1}},
            {"$limit": 100},
        ]
        return list(db["students"].aggregate(pipeline, maxTimeMS=5000))
    except Exception as exc:
        logger.warning(f"MongoDB aggregation fallback triggered: {exc}")
        return None


def compute_advanced_analytics(user, raw_filters=None):
    """
    Compute multi-dimensional academic analytics strictly scoped to `user` RBAC permissions:
      - Student enrollment statistics
      - Branch-, semester-, academic-year-, and course-wise breakdowns
      - Attendance summaries and chronological date trends
      - Grade distributions (using active GradingScale) and course performance comparisons
      - Credit-weighted GPA/CGPA summaries
    """
    val_res = validate_and_normalize_filters(
        {**(raw_filters or {}), "report_type": "student_directory", "format": "csv"}
    )
    if "error" in val_res:
        return {"error": val_res["error"], "status_code": 400}

    rbac_res = enforce_report_rbac(user, val_res)
    if "error" in rbac_res:
        return rbac_res

    filters = rbac_res["filters"]
    scale = get_grading_scale()

    # 1. Load RBAC-scoped students and apply multi-field filters
    all_students = view_students(
        branch=filters["branch"] or None,
        status=filters["status"] or None,
        semester=filters["semester"] or None,
    )
    scoped_students = filter_students_for_user(user, all_students)

    if filters["reg_no"]:
        scoped_students = [
            s for s in scoped_students if s.get("reg_no") == filters["reg_no"]
        ]
    if filters["academic_year"]:
        scoped_students = [
            s
            for s in scoped_students
            if str(s.get("academic_year") or "2025-2026") == filters["academic_year"]
        ]

    allowed_reg_nos = {s["reg_no"] for s in scoped_students}

    # 2. Load RBAC-scoped courses, enrollments, attendance, and grades
    courses = list_courses(
        user=user,
        branch=filters["branch"] or None,
        semester=filters["semester"] or None,
        include_archived=False,
    )
    if filters["course_code"]:
        courses = [
            c for c in courses if c.get("course_code") == filters["course_code"]
        ]
    if filters["academic_year"]:
        courses = [
            c
            for c in courses
            if str(c.get("academic_year") or "2025-2026") == filters["academic_year"]
        ]

    enrollments = list_enrollments(
        user=user,
        reg_no=filters["reg_no"] or None,
        course_code=filters["course_code"] or None,
        semester=filters["semester"] or None,
    )
    if filters["branch"]:
        enrollments = [
            e for e in enrollments if str(e.get("branch") or "").upper() == filters["branch"]
        ]
    if filters["academic_year"]:
        enrollments = [
            e
            for e in enrollments
            if str(e.get("academic_year") or "2025-2026") == filters["academic_year"]
        ]

    # If a course filter is active, narrow students to those enrolled in the course
    if filters["course_code"]:
        enrolled_regs = {e["reg_no"] for e in enrollments}
        scoped_students = [
            s for s in scoped_students if s["reg_no"] in enrolled_regs
        ]
        allowed_reg_nos = {s["reg_no"] for s in scoped_students}

    attendance_records = list_attendance(
        user=user,
        reg_no=filters["reg_no"] or None,
        course_code=filters["course_code"] or None,
    )
    if filters["branch"]:
        attendance_records = [
            a
            for a in attendance_records
            if str(a.get("branch") or "").upper() == filters["branch"]
        ]
    if filters["date_from"]:
        attendance_records = [
            a for a in attendance_records if str(a.get("date") or "") >= filters["date_from"]
        ]
    if filters["date_to"]:
        attendance_records = [
            a for a in attendance_records if str(a.get("date") or "") <= filters["date_to"]
        ]
    if filters["semester"] or filters["academic_year"] or filters["status"]:
        attendance_records = [
            a for a in attendance_records if a.get("reg_no") in allowed_reg_nos
        ]

    grades_records = list_grades(
        user=user,
        reg_no=filters["reg_no"] or None,
        course_code=filters["course_code"] or None,
        semester=filters["semester"] or None,
    )
    if filters["branch"]:
        grades_records = [
            g for g in grades_records if str(g.get("branch") or "").upper() == filters["branch"]
        ]
    if filters["academic_year"]:
        grades_records = [
            g
            for g in grades_records
            if str(g.get("academic_year") or "2025-2026") == filters["academic_year"]
        ]
    if filters["status"]:
        grades_records = [
            g for g in grades_records if g.get("reg_no") in allowed_reg_nos
        ]

    # 3. Branch, Semester, Academic-Year & Course Enrollment Breakdowns
    branch_buckets = {}
    semester_buckets = {}
    year_buckets = {}

    for s in scoped_students:
        b = str(s.get("branch") or "Unassigned").strip().upper()
        sem = str(s.get("semester") or "1").strip()
        yr = str(s.get("academic_year") or "2025-2026").strip()

        b_entry = branch_buckets.setdefault(
            b, {"count": 0, "gpas": [], "att": [], "ages": []}
        )
        b_entry["count"] += 1
        if s.get("age") is not None:
            b_entry["ages"].append(float(s["age"]))
        if s.get("gpa") is not None:
            b_entry["gpas"].append(float(s["gpa"]))
        if s.get("attendance_pct") is not None:
            b_entry["att"].append(float(s["attendance_pct"]))

        semester_buckets[sem] = semester_buckets.get(sem, 0) + 1
        year_buckets[yr] = year_buckets.get(yr, 0) + 1

    # Attempt MongoDB aggregation pipeline when available for Admin un-narrowed queries
    if user.get("role") == "Admin" and not filters["course_code"]:
        mongo_match = {}
        if filters["branch"]:
            mongo_match["branch"] = filters["branch"]
        if filters["semester"]:
            mongo_match["semester"] = filters["semester"]
        if filters["status"]:
            mongo_match["status"] = filters["status"]
        _run_mongo_branch_pipeline(mongo_match)

    enrollment_by_branch = []
    for b, data in sorted(
        branch_buckets.items(), key=lambda x: (-x[1]["count"], x[0])
    ):
        enrollment_by_branch.append({
            "branch": b,
            "student_count": data["count"],
            "avg_age": round(sum(data["ages"]) / len(data["ages"]), 1)
            if data["ages"]
            else 0.0,
            "avg_gpa": round(sum(data["gpas"]) / len(data["gpas"]), 2)
            if data["gpas"]
            else None,
            "avg_attendance": round(sum(data["att"]) / len(data["att"]), 1)
            if data["att"]
            else None,
        })

    enrollment_by_semester = [
        {"semester": f"Sem {sem}", "raw_semester": sem, "student_count": cnt}
        for sem, cnt in sorted(
            semester_buckets.items(),
            key=lambda x: (int(x[0]) if str(x[0]).isdigit() else 99, str(x[0])),
        )
    ]

    enrollment_by_academic_year = [
        {"academic_year": yr, "student_count": cnt}
        for yr, cnt in sorted(year_buckets.items())
    ]

    enrollment_by_course = []
    for c in courses:
        cap = max(1, int(c.get("capacity") or 60))
        enr_cnt = int(c.get("enrolled_count") or 0)
        enrollment_by_course.append({
            "course_code": c["course_code"],
            "course_name": c["course_name"],
            "branch": c["branch"],
            "semester": c.get("semester", "1"),
            "academic_year": c.get("academic_year", "2025-2026"),
            "credits": c.get("credits", 3.0),
            "enrolled_count": enr_cnt,
            "capacity": cap,
            "utilization_pct": round((enr_cnt / cap) * 100.0, 1),
        })

    # 4. Attendance Summaries & Chronological Date Trends
    att_metrics = compute_attendance_metrics(
        attendance_records, scale.late_attendance_weight
    )
    by_student_att = {}
    by_date_att = {}
    by_course_att = {}

    for rec in attendance_records:
        by_student_att.setdefault(rec["reg_no"], []).append(rec)
        by_date_att.setdefault(rec["date"], []).append(rec)
        by_course_att.setdefault(rec["course_code"], []).append(rec)

    below_threshold_count = 0
    for rn, recs in by_student_att.items():
        m = compute_attendance_metrics(recs, scale.late_attendance_weight)
        if m["attendance_percentage"] < scale.min_attendance_threshold:
            below_threshold_count += 1

    # Also check student profile attendance_pct if no session records exist yet
    if not by_student_att:
        for s in scoped_students:
            if (
                s.get("attendance_pct") is not None
                and float(s["attendance_pct"]) < scale.min_attendance_threshold
            ):
                below_threshold_count += 1

    attendance_trends = []
    for dt in sorted(by_date_att.keys()):
        dm = compute_attendance_metrics(
            by_date_att[dt], scale.late_attendance_weight
        )
        attendance_trends.append({
            "date": dt,
            "present": dm["present"],
            "absent": dm["absent"],
            "late": dm["late"],
            "excused": dm["excused"],
            "total_sessions": dm["total_sessions"],
            "attendance_percentage": dm["attendance_percentage"],
        })

    attendance_summary = {
        **att_metrics,
        "below_threshold_count": below_threshold_count,
        "min_attendance_threshold": scale.min_attendance_threshold,
        "formula": "Attendance % = ((Present + 0.5 * Late) / (Present + Absent + Late)) * 100 (Excused excluded)",
    }

    # 5. Grade Distributions & Course Performance Comparisons
    band_counts = {b["letter_grade"]: 0 for b in scale.bands}
    for g in grades_records:
        lg = g.get("letter_grade")
        if lg in band_counts:
            band_counts[lg] += 1
        else:
            band_counts[lg] = band_counts.get(lg, 0) + 1

    total_grades_count = len(grades_records)
    grade_distribution = []
    for b in scale.bands:
        lg = b["letter_grade"]
        cnt = band_counts.get(lg, 0)
        grade_distribution.append({
            "letter_grade": lg,
            "grade_points": b["grade_points"],
            "min_percentage": b["min_percentage"],
            "count": cnt,
            "share_pct": round((cnt / total_grades_count) * 100.0, 1)
            if total_grades_count > 0
            else 0.0,
        })

    by_course_grades = {}
    for g in grades_records:
        by_course_grades.setdefault(g["course_code"], []).append(g)

    all_course_codes = sorted(
        {c["course_code"] for c in courses}
        | set(by_course_grades.keys())
        | set(by_course_att.keys())
    )

    course_performance = []
    for cc in all_course_codes:
        c_doc = get_course(cc) or {}
        if user.get("role") == "Faculty" and c_doc and not can_manage_course(user, c_doc):
            continue
        c_grades = by_course_grades.get(cc, [])
        c_att_recs = by_course_att.get(cc, [])
        c_att_metrics = compute_attendance_metrics(
            c_att_recs, scale.late_attendance_weight
        )

        pcts = [float(g.get("percentage", 0.0)) for g in c_grades]
        gps = [float(g.get("grade_points", 0.0)) for g in c_grades]
        pass_cnt = sum(1 for g in c_grades if g.get("letter_grade") != "F")

        course_performance.append({
            "course_code": cc,
            "course_name": c_doc.get("course_name", cc),
            "branch": c_doc.get("branch", ""),
            "semester": str(c_doc.get("semester", "1")),
            "credits": float(c_doc.get("credits", 3.0)),
            "enrolled_count": sum(
                1 for e in enrollments if e.get("course_code") == cc
            ),
            "assessments_count": len(c_grades),
            "average_percentage": round(sum(pcts) / len(pcts), 2)
            if pcts
            else None,
            "average_grade_points": round(sum(gps) / len(gps), 2)
            if gps
            else None,
            "pass_rate": round((pass_cnt / len(c_grades)) * 100.0, 1)
            if c_grades
            else None,
            "attendance_percentage": c_att_metrics["attendance_percentage"]
            if c_att_recs
            else None,
            "total_attendance_entries": c_att_metrics["total_sessions"],
        })

    # 6. Credit-Weighted GPA / CGPA Summary across Scoped Cohort
    cgpa_values = []
    for s in scoped_students:
        calc = calculate_student_gpa(
            s["reg_no"], semester=filters["semester"] or None
        )
        if calc.get("cgpa") is not None:
            cgpa_values.append(float(calc["cgpa"]))
        elif s.get("gpa") is not None:
            cgpa_values.append(float(s["gpa"]))

    max_gp = float(scale.max_grade_points or 10.0)
    distinction_cutoff = 0.85 * max_gp
    support_cutoff = 0.55 * max_gp

    gpa_summary = {
        "scale_name": scale.scale_name,
        "max_grade_points": max_gp,
        "formula": "CGPA = sum(Course Credits * Grade Points) / sum(Course Credits)",
        "evaluated_students": len(cgpa_values),
        "average_cgpa": round(sum(cgpa_values) / len(cgpa_values), 2)
        if cgpa_values
        else None,
        "highest_cgpa": round(max(cgpa_values), 2) if cgpa_values else None,
        "lowest_cgpa": round(min(cgpa_values), 2) if cgpa_values else None,
        "distinction_count": sum(
            1 for v in cgpa_values if v >= distinction_cutoff
        ),
        "support_needed_count": sum(
            1 for v in cgpa_values if v < support_cutoff
        ),
    }

    return {
        "success": True,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "role_scope": user.get("role"),
        "applied_filters": {
            "branch": filters["branch"] or "All",
            "semester": filters["semester"] or "All",
            "academic_year": filters["academic_year"] or "All",
            "course_code": filters["course_code"] or "All",
            "reg_no": filters["reg_no"] or "All",
            "status": filters["status"] or "All",
            "date_from": filters["date_from"] or None,
            "date_to": filters["date_to"] or None,
        },
        "totals": {
            "total_students": len(scoped_students),
            "total_courses": len(courses),
            "total_enrollments": len(enrollments),
            "total_attendance_records": len(attendance_records),
            "total_grade_records": len(grades_records),
        },
        "enrollment_by_branch": enrollment_by_branch,
        "enrollment_by_semester": enrollment_by_semester,
        "enrollment_by_academic_year": enrollment_by_academic_year,
        "enrollment_by_course": enrollment_by_course,
        "attendance_summary": attendance_summary,
        "attendance_trends": attendance_trends,
        "grade_distribution": grade_distribution,
        "course_performance": course_performance,
        "gpa_summary": gpa_summary,
    }


# ==========================================
# OBJECTIVE B: STRUCTURED REPORT BUILDER & EXPORTS (CSV, EXCEL, PDF)
# ==========================================
def build_report_dataset(user, raw_filters=None):
    """
    Validate filters, enforce RBAC, and construct a structured report dataset:
      {
        "report_type": ...,
        "title": ...,
        "generated_at": ...,
        "generated_by": ...,
        "applied_filters": ...,
        "columns": [{"key": ..., "label": ...}],
        "rows": [...],
        "summary_metrics": [...],
        "secondary_sheets": [...] (optional extra worksheets for Excel / tables for PDF)
      }
    """
    val_res = validate_and_normalize_filters(raw_filters)
    if "error" in val_res:
        return {"error": val_res["error"], "status_code": 400}

    rbac_res = enforce_report_rbac(user, val_res)
    if "error" in rbac_res:
        return rbac_res

    filters = rbac_res["filters"]
    rtype = filters["report_type"]
    limit = filters["limit"]
    scale = get_grading_scale()

    analytics = compute_advanced_analytics(user, filters)
    if "error" in analytics:
        return analytics

    columns = []
    rows = []
    summary_metrics = []
    secondary_sheets = []

    if rtype == "student_directory":
        all_students = view_students(
            branch=filters["branch"] or None,
            status=filters["status"] or None,
            semester=filters["semester"] or None,
        )
        scoped = filter_students_for_user(user, all_students)
        if filters["reg_no"]:
            scoped = [s for s in scoped if s.get("reg_no") == filters["reg_no"]]
        if filters["academic_year"]:
            scoped = [
                s
                for s in scoped
                if str(s.get("academic_year") or "2025-2026")
                == filters["academic_year"]
            ]

        columns = [
            {"key": "reg_no", "label": "Registration No"},
            {"key": "roll_no", "label": "Roll No"},
            {"key": "name", "label": "Student Name"},
            {"key": "branch", "label": "Branch"},
            {"key": "semester", "label": "Semester"},
            {"key": "academic_year", "label": "Academic Year"},
            {"key": "age", "label": "Age"},
            {"key": "status", "label": "Status"},
            {"key": "attendance_pct", "label": "Attendance (%)"},
            {"key": "gpa", "label": "CGPA"},
        ]
        for s in scoped[:limit]:
            rows.append({
                "reg_no": s.get("reg_no", ""),
                "roll_no": s.get("roll_no", ""),
                "name": s.get("name", ""),
                "branch": s.get("branch", ""),
                "semester": s.get("semester", "1"),
                "academic_year": s.get("academic_year", "2025-2026"),
                "age": s.get("age", ""),
                "status": s.get("status", "Active"),
                "attendance_pct": s.get("attendance_pct")
                if s.get("attendance_pct") is not None
                else "N/A",
                "gpa": s.get("gpa") if s.get("gpa") is not None else "N/A",
            })
        summary_metrics = [
            {"label": "Total Students", "value": len(rows)},
            {
                "label": "Branches Covered",
                "value": len(analytics["enrollment_by_branch"]),
            },
            {
                "label": "Cohort Avg CGPA",
                "value": analytics["gpa_summary"]["average_cgpa"]
                if analytics["gpa_summary"]["average_cgpa"] is not None
                else "N/A",
            },
        ]

    elif rtype == "course_enrollment":
        enrollments = list_enrollments(
            user=user,
            reg_no=filters["reg_no"] or None,
            course_code=filters["course_code"] or None,
            semester=filters["semester"] or None,
        )
        if filters["branch"]:
            enrollments = [
                e
                for e in enrollments
                if str(e.get("branch") or "").upper() == filters["branch"]
            ]
        if filters["academic_year"]:
            enrollments = [
                e
                for e in enrollments
                if str(e.get("academic_year") or "2025-2026")
                == filters["academic_year"]
            ]

        columns = [
            {"key": "enrollment_id", "label": "Enrollment ID"},
            {"key": "reg_no", "label": "Registration No"},
            {"key": "student_name", "label": "Student Name"},
            {"key": "course_code", "label": "Course Code"},
            {"key": "course_name", "label": "Course Name"},
            {"key": "branch", "label": "Branch"},
            {"key": "credits", "label": "Credits"},
            {"key": "semester", "label": "Semester"},
            {"key": "academic_year", "label": "Academic Year"},
            {"key": "status", "label": "Status"},
        ]
        for e in enrollments[:limit]:
            rows.append({
                "enrollment_id": e.get("enrollment_id", ""),
                "reg_no": e.get("reg_no", ""),
                "student_name": e.get("student_name", ""),
                "course_code": e.get("course_code", ""),
                "course_name": e.get("course_name", ""),
                "branch": e.get("branch", ""),
                "credits": e.get("credits", 3.0),
                "semester": e.get("semester", "1"),
                "academic_year": e.get("academic_year", "2025-2026"),
                "status": e.get("status", "Enrolled"),
            })
        summary_metrics = [
            {"label": "Total Enrollments", "value": len(rows)},
            {
                "label": "Active Courses",
                "value": len(analytics["enrollment_by_course"]),
            },
        ]
        secondary_sheets.append({
            "title": "Course Utilization",
            "columns": [
                {"key": "course_code", "label": "Course Code"},
                {"key": "course_name", "label": "Course Name"},
                {"key": "branch", "label": "Branch"},
                {"key": "credits", "label": "Credits"},
                {"key": "enrolled_count", "label": "Enrolled"},
                {"key": "capacity", "label": "Capacity"},
                {"key": "utilization_pct", "label": "Utilization (%)"},
            ],
            "rows": analytics["enrollment_by_course"],
        })

    elif rtype == "attendance_register":
        recs = list_attendance(
            user=user,
            reg_no=filters["reg_no"] or None,
            course_code=filters["course_code"] or None,
        )
        if filters["branch"]:
            recs = [
                r
                for r in recs
                if str(r.get("branch") or "").upper() == filters["branch"]
            ]
        if filters["date_from"]:
            recs = [
                r for r in recs if str(r.get("date") or "") >= filters["date_from"]
            ]
        if filters["date_to"]:
            recs = [
                r for r in recs if str(r.get("date") or "") <= filters["date_to"]
            ]

        columns = [
            {"key": "date", "label": "Session Date"},
            {"key": "session", "label": "Session"},
            {"key": "course_code", "label": "Course Code"},
            {"key": "course_name", "label": "Course Name"},
            {"key": "reg_no", "label": "Registration No"},
            {"key": "student_name", "label": "Student Name"},
            {"key": "branch", "label": "Branch"},
            {"key": "status", "label": "Status"},
            {"key": "remarks", "label": "Remarks"},
        ]
        for r in recs[:limit]:
            rows.append({
                "date": r.get("date", ""),
                "session": r.get("session", "Lecture"),
                "course_code": r.get("course_code", ""),
                "course_name": r.get("course_name", ""),
                "reg_no": r.get("reg_no", ""),
                "student_name": r.get("student_name", ""),
                "branch": r.get("branch", ""),
                "status": r.get("status", "Present"),
                "remarks": r.get("remarks", ""),
            })
        att_sum = analytics["attendance_summary"]
        summary_metrics = [
            {"label": "Total Sessions Logged", "value": att_sum["total_sessions"]},
            {
                "label": "Overall Attendance (%)",
                "value": f"{att_sum['attendance_percentage']}%",
            },
            {
                "label": "Below Threshold (<75%)",
                "value": att_sum["below_threshold_count"],
            },
        ]

    elif rtype == "grade_sheet":
        grades = list_grades(
            user=user,
            reg_no=filters["reg_no"] or None,
            course_code=filters["course_code"] or None,
            semester=filters["semester"] or None,
        )
        if filters["branch"]:
            grades = [
                g
                for g in grades
                if str(g.get("branch") or "").upper() == filters["branch"]
            ]
        if filters["academic_year"]:
            grades = [
                g
                for g in grades
                if str(g.get("academic_year") or "2025-2026")
                == filters["academic_year"]
            ]

        columns = [
            {"key": "reg_no", "label": "Registration No"},
            {"key": "student_name", "label": "Student Name"},
            {"key": "course_code", "label": "Course Code"},
            {"key": "course_name", "label": "Course Name"},
            {"key": "branch", "label": "Branch"},
            {"key": "semester", "label": "Semester"},
            {"key": "assessment_type", "label": "Assessment"},
            {"key": "marks_obtained", "label": "Marks Obtained"},
            {"key": "max_marks", "label": "Max Marks"},
            {"key": "percentage", "label": "Percentage (%)"},
            {"key": "letter_grade", "label": "Letter Grade"},
            {"key": "grade_points", "label": "Grade Points"},
            {"key": "credits", "label": "Credits"},
        ]
        for g in grades[:limit]:
            rows.append({
                "reg_no": g.get("reg_no", ""),
                "student_name": g.get("student_name", ""),
                "course_code": g.get("course_code", ""),
                "course_name": g.get("course_name", ""),
                "branch": g.get("branch", ""),
                "semester": g.get("semester", "1"),
                "assessment_type": g.get("assessment_type", "Overall"),
                "marks_obtained": g.get("marks_obtained", 0),
                "max_marks": g.get("max_marks", 100),
                "percentage": g.get("percentage", 0),
                "letter_grade": g.get("letter_grade", ""),
                "grade_points": g.get("grade_points", 0),
                "credits": g.get("credits", 3.0),
            })
        summary_metrics = [
            {"label": "Assessments Graded", "value": len(rows)},
            {"label": "Grading Scale", "value": scale.scale_name},
            {
                "label": "Cohort Avg CGPA",
                "value": analytics["gpa_summary"]["average_cgpa"]
                if analytics["gpa_summary"]["average_cgpa"] is not None
                else "N/A",
            },
        ]
        secondary_sheets.append({
            "title": "Grade Distribution",
            "columns": [
                {"key": "letter_grade", "label": "Letter Grade"},
                {"key": "min_percentage", "label": "Min Percentage (%)"},
                {"key": "grade_points", "label": "Grade Points"},
                {"key": "count", "label": "Count"},
                {"key": "share_pct", "label": "Share (%)"},
            ],
            "rows": analytics["grade_distribution"],
        })

    elif rtype == "academic_transcript":
        all_students = view_students(
            branch=filters["branch"] or None,
            semester=filters["semester"] or None,
        )
        scoped = filter_students_for_user(user, all_students)
        if filters["reg_no"]:
            scoped = [s for s in scoped if s.get("reg_no") == filters["reg_no"]]

        columns = [
            {"key": "reg_no", "label": "Registration No"},
            {"key": "student_name", "label": "Student Name"},
            {"key": "branch", "label": "Branch"},
            {"key": "course_code", "label": "Course Code"},
            {"key": "course_name", "label": "Course Name"},
            {"key": "semester", "label": "Semester"},
            {"key": "academic_year", "label": "Academic Year"},
            {"key": "credits", "label": "Credits"},
            {"key": "percentage", "label": "Course (%)"},
            {"key": "letter_grade", "label": "Letter Grade"},
            {"key": "grade_points", "label": "Grade Points"},
            {"key": "student_cgpa", "label": "Cumulative CGPA"},
        ]

        for s in scoped:
            gpa_data = calculate_student_gpa(
                s["reg_no"], semester=filters["semester"] or None
            )
            cgpa_val = (
                gpa_data["cgpa"]
                if gpa_data["cgpa"] is not None
                else (s.get("gpa") if s.get("gpa") is not None else "N/A")
            )
            if gpa_data["courses"]:
                for c_item in gpa_data["courses"]:
                    if (
                        filters["course_code"]
                        and c_item["course_code"] != filters["course_code"]
                    ):
                        continue
                    rows.append({
                        "reg_no": s["reg_no"],
                        "student_name": s.get("name", ""),
                        "branch": s.get("branch", ""),
                        "course_code": c_item["course_code"],
                        "course_name": c_item["course_name"],
                        "semester": c_item["semester"],
                        "academic_year": c_item["academic_year"],
                        "credits": c_item["credits"],
                        "percentage": c_item["percentage"],
                        "letter_grade": c_item["letter_grade"],
                        "grade_points": c_item["grade_points"],
                        "student_cgpa": cgpa_val,
                    })
                    if len(rows) >= limit:
                        break
            else:
                if not filters["course_code"]:
                    rows.append({
                        "reg_no": s["reg_no"],
                        "student_name": s.get("name", ""),
                        "branch": s.get("branch", ""),
                        "course_code": "N/A",
                        "course_name": "No graded courses yet",
                        "semester": s.get("semester", "1"),
                        "academic_year": s.get("academic_year", "2025-2026"),
                        "credits": 0,
                        "percentage": "N/A",
                        "letter_grade": "N/A",
                        "grade_points": "N/A",
                        "student_cgpa": cgpa_val,
                    })
            if len(rows) >= limit:
                break

        summary_metrics = [
            {"label": "Transcript Entries", "value": len(rows)},
            {"label": "Grading Scale", "value": scale.scale_name},
            {
                "label": "Average CGPA",
                "value": analytics["gpa_summary"]["average_cgpa"]
                if analytics["gpa_summary"]["average_cgpa"] is not None
                else "N/A",
            },
        ]

    elif rtype == "executive_summary":
        columns = [
            {"key": "branch", "label": "Branch / Department"},
            {"key": "student_count", "label": "Enrolled Students"},
            {"key": "avg_age", "label": "Average Age (yrs)"},
            {"key": "avg_gpa", "label": "Average CGPA"},
            {"key": "avg_attendance", "label": "Average Attendance (%)"},
        ]
        for b_item in analytics["enrollment_by_branch"][:limit]:
            rows.append({
                "branch": b_item["branch"],
                "student_count": b_item["student_count"],
                "avg_age": b_item["avg_age"],
                "avg_gpa": b_item["avg_gpa"]
                if b_item["avg_gpa"] is not None
                else "N/A",
                "avg_attendance": b_item["avg_attendance"]
                if b_item["avg_attendance"] is not None
                else "N/A",
            })

        summary_metrics = [
            {
                "label": "Total Students",
                "value": analytics["totals"]["total_students"],
            },
            {
                "label": "Active Courses",
                "value": analytics["totals"]["total_courses"],
            },
            {
                "label": "Overall Attendance (%)",
                "value": f"{analytics['attendance_summary']['attendance_percentage']}%",
            },
            {
                "label": "Cohort Avg CGPA",
                "value": analytics["gpa_summary"]["average_cgpa"]
                if analytics["gpa_summary"]["average_cgpa"] is not None
                else "N/A",
            },
        ]

        secondary_sheets.extend([
            {
                "title": "Semester Breakdown",
                "columns": [
                    {"key": "semester", "label": "Semester"},
                    {"key": "student_count", "label": "Student Count"},
                ],
                "rows": analytics["enrollment_by_semester"],
            },
            {
                "title": "Course Performance",
                "columns": [
                    {"key": "course_code", "label": "Course Code"},
                    {"key": "course_name", "label": "Course Name"},
                    {"key": "branch", "label": "Branch"},
                    {"key": "enrolled_count", "label": "Enrolled"},
                    {"key": "average_percentage", "label": "Avg Marks (%)"},
                    {"key": "pass_rate", "label": "Pass Rate (%)"},
                    {"key": "attendance_percentage", "label": "Attendance (%)"},
                ],
                "rows": analytics["course_performance"],
            },
        ])

    generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    return {
        "success": True,
        "report_type": rtype,
        "format": filters["format"],
        "title": VALID_REPORT_TYPES[rtype],
        "generated_at": generated_at,
        "generated_by": f"{user.get('username', 'user')} ({user.get('role', 'User')})",
        "applied_filters": analytics["applied_filters"],
        "columns": columns,
        "rows": rows,
        "row_count": len(rows),
        "summary_metrics": summary_metrics,
        "secondary_sheets": secondary_sheets,
    }


def generate_csv_bytes(report_data):
    """Generate formula-sanitized UTF-8 CSV bytes with metadata header."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")

    # Metadata block
    writer.writerow([
        sanitize_spreadsheet_cell(f"Report: {report_data['title']}"),
        sanitize_spreadsheet_cell(f"Generated At: {report_data['generated_at']}"),
        sanitize_spreadsheet_cell(f"Scope: {report_data['generated_by']}"),
    ])
    af = report_data.get("applied_filters") or {}
    filter_desc = ", ".join(f"{k}={v}" for k, v in af.items() if v)
    writer.writerow([sanitize_spreadsheet_cell(f"Filters: {filter_desc}")])
    writer.writerow([])

    # Column headers
    cols = report_data.get("columns") or []
    writer.writerow([sanitize_spreadsheet_cell(c["label"]) for c in cols])

    # Data rows
    for row in report_data.get("rows") or []:
        writer.writerow([
            sanitize_spreadsheet_cell(row.get(c["key"], "")) for c in cols
        ])

    return buf.getvalue().encode("utf-8-sig")


def _populate_excel_sheet(ws, title, columns, rows, metadata_lines=None):
    header_fill = PatternFill(
        start_color="1E3A8A", end_color="1E3A8A", fill_type="solid"
    )
    header_font = Font(color="FFFFFF", bold=True, size=11)
    title_font = Font(color="0F172A", bold=True, size=14)
    meta_font = Font(color="475569", italic=True, size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    current_row = 1
    ws.cell(row=current_row, column=1, value=sanitize_spreadsheet_cell(title)).font = (
        title_font
    )
    current_row += 1

    for line in metadata_lines or []:
        ws.cell(
            row=current_row, column=1, value=sanitize_spreadsheet_cell(line)
        ).font = meta_font
        current_row += 1

    current_row += 1  # Spacer row

    # Header row
    for col_idx, col_def in enumerate(columns, start=1):
        cell = ws.cell(
            row=current_row,
            column=col_idx,
            value=sanitize_spreadsheet_cell(col_def["label"]),
        )
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = thin_border

    current_row += 1

    # Data rows
    if not rows:
        ws.cell(
            row=current_row,
            column=1,
            value="No matching records found for the selected filters.",
        )
    else:
        for r_item in rows:
            for col_idx, col_def in enumerate(columns, start=1):
                raw_val = r_item.get(col_def["key"], "")
                safe_val = sanitize_spreadsheet_cell(raw_val)
                cell = ws.cell(row=current_row, column=col_idx, value=safe_val)
                cell.border = thin_border
                if isinstance(safe_val, float):
                    cell.number_format = "0.00"
            current_row += 1

    # Auto-fit sensible column widths (min 12, max 42)
    for col_idx, col_def in enumerate(columns, start=1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(col_def["label"]))
        for r_item in rows[:200]:
            val_str = str(r_item.get(col_def["key"], ""))
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(12, min(42, max_len + 4))


def generate_excel_bytes(report_data):
    """Generate a multi-worksheet Excel (.xlsx) workbook using openpyxl with formula-injection protection."""
    wb = Workbook()
    ws_main = wb.active
    ws_main.title = "Report Data"

    af = report_data.get("applied_filters") or {}
    filter_summary = ", ".join(f"{k}={v}" for k, v in af.items() if v)
    meta_lines = [
        f"Generated At: {report_data['generated_at']} | Generated By: {report_data['generated_by']}",
        f"Applied Filters: {filter_summary}",
    ]

    _populate_excel_sheet(
        ws_main,
        report_data["title"],
        report_data.get("columns") or [],
        report_data.get("rows") or [],
        metadata_lines=meta_lines,
    )

    # Summary & Metadata Worksheet
    ws_meta = wb.create_sheet(title="Summary & Metadata")
    _populate_excel_sheet(
        ws_meta,
        f"{report_data['title']} — Key Summary Metrics",
        [
            {"key": "label", "label": "Metric"},
            {"key": "value", "label": "Value"},
        ],
        report_data.get("summary_metrics") or [],
        metadata_lines=meta_lines,
    )

    # Any secondary breakdown sheets
    for extra in report_data.get("secondary_sheets") or []:
        sheet_title = re.sub(r"[\\/*?:\[\]]", "", str(extra["title"]))[:31]
        ws_extra = wb.create_sheet(title=sheet_title or "Breakdown")
        _populate_excel_sheet(
            ws_extra,
            extra["title"],
            extra.get("columns") or [],
            extra.get("rows") or [],
            metadata_lines=meta_lines,
        )

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_pdf_bytes(report_data):
    """Generate a formatted PDF document with tables, metadata header, and page numbers via ReportLab."""
    buffer = io.BytesIO()
    cols = report_data.get("columns") or []
    page_size = landscape(A4) if len(cols) > 6 else A4

    doc = SimpleDocTemplate(
        buffer,
        pagesize=page_size,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=report_data["title"],
        author="Student Database Management System 2.0",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=15,
        leading=19,
        textColor=colors.HexColor("#1E3A8A"),
        spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        "ReportMeta",
        parent=styles["Normal"],
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor("#475569"),
        spaceAfter=6,
    )
    cell_style = ParagraphStyle(
        "TableCell",
        parent=styles["Normal"],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#0F172A"),
    )
    header_cell_style = ParagraphStyle(
        "TableHeaderCell",
        parent=styles["Normal"],
        fontSize=8,
        leading=10,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )

    elements = []
    elements.append(Paragraph(report_data["title"], title_style))

    af = report_data.get("applied_filters") or {}
    filter_summary = ", ".join(f"{k}: {v}" for k, v in af.items() if v)
    elements.append(
        Paragraph(
            f"<b>Generated At:</b> {report_data['generated_at']} &nbsp;|&nbsp; "
            f"<b>Generated By:</b> {report_data['generated_by']}<br/>"
            f"<b>Applied Filters:</b> {filter_summary}",
            meta_style,
        )
    )

    # Summary KPI strip
    summary_metrics = report_data.get("summary_metrics") or []
    if summary_metrics:
        kpi_headers = [m["label"] for m in summary_metrics]
        kpi_vals = [str(m["value"]) for m in summary_metrics]
        kpi_table = Table([kpi_headers, kpi_vals])
        kpi_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EFF6FF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#BFDBFE")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
            ])
        )
        elements.append(kpi_table)
        elements.append(Spacer(1, 5 * mm))

    # Main Data Table
    rows = report_data.get("rows") or []
    if not rows:
        elements.append(
            Paragraph(
                "<i>No records matched the selected filter criteria.</i>",
                styles["Normal"],
            )
        )
    else:
        table_rows = [
            [Paragraph(str(c["label"]), header_cell_style) for c in cols]
        ]
        for r_item in rows:
            table_rows.append([
                Paragraph(
                    str(r_item.get(c["key"], ""))
                    .replace("&", "&amp;")
                    .replace("<", "&lt;")
                    .replace(">", "&gt;"),
                    cell_style,
                )
                for c in cols
            ])

        main_table = Table(table_rows, repeatRows=1)
        main_table.setStyle(
            TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E3A8A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F8FAFC")],
                ),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )
        elements.append(main_table)

    def _draw_header_footer(canvas_obj, doc_obj):
        canvas_obj.saveState()
        canvas_obj.setFont("Helvetica", 8)
        canvas_obj.setFillColor(colors.HexColor("#64748B"))
        width, height = page_size
        canvas_obj.drawString(
            14 * mm,
            10 * mm,
            "Student Database Management System 2.0 — Official Academic Report",
        )
        canvas_obj.drawRightString(
            width - 14 * mm,
            10 * mm,
            f"Page {doc_obj.page}",
        )
        canvas_obj.restoreState()

    doc.build(
        elements,
        onFirstPage=_draw_header_footer,
        onLaterPages=_draw_header_footer,
    )
    return buffer.getvalue()


def export_report_file(user, raw_filters=None):
    """
    Validate filters, enforce RBAC, build the report dataset, and render the requested
    export artifact (CSV, Excel .xlsx, or PDF) with a safe filename and MIME type.
    """
    dataset = build_report_dataset(user, raw_filters)
    if "error" in dataset:
        return dataset

    fmt = dataset["format"]
    suffix = (
        dataset["applied_filters"].get("course_code")
        if dataset["applied_filters"].get("course_code") != "All"
        else dataset["applied_filters"].get("branch", "")
    )
    if suffix == "All":
        suffix = ""

    filename = build_safe_filename(dataset["report_type"], fmt, suffix=suffix)

    if fmt == "xlsx":
        content = generate_excel_bytes(dataset)
        mimetype = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    elif fmt == "pdf":
        content = generate_pdf_bytes(dataset)
        mimetype = "application/pdf"
    else:
        content = generate_csv_bytes(dataset)
        mimetype = "text/csv; charset=utf-8"

    return {
        "success": True,
        "filename": filename,
        "mimetype": mimetype,
        "content": content,
        "report_type": dataset["report_type"],
        "row_count": dataset["row_count"],
        "applied_filters": dataset["applied_filters"],
    }

