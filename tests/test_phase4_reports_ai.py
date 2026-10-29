import csv
import io
import os
import urllib.error

os.environ["STUDENTDB_TEST_MODE"] = "1"
os.environ["SECRET_KEY"] = "test-secret-key-for-pytest-suite-2026"

from openpyxl import load_workbook
import pytest

from backend.api.app import app, clear_rate_limits
from services.academic_service import (
    create_course,
    enroll_student,
    record_attendance,
    record_grade,
    reset_test_academics,
)
from services.ai_assistant_service import set_custom_llm_caller
from services.auth_service import reset_test_auth
from services.report_service import (
    build_safe_filename,
    sanitize_spreadsheet_cell,
)
from services.student_service import add_student, reset_test_students


@pytest.fixture(autouse=True)
def isolated_phase4_fixture():
    os.environ["STUDENTDB_TEST_MODE"] = "1"
    os.environ["AI_ASSISTANT_PROVIDER"] = "deterministic"
    os.environ.pop("AI_ASSISTANT_API_KEY", None)
    clear_rate_limits()
    set_custom_llm_caller(None)
    reset_test_students()
    reset_test_auth()
    reset_test_academics()

    # Seed known dataset
    add_student(
        "REG2026001",
        "101",
        "Aarav Sharma",
        20,
        "CSE",
        email="aarav@univ.edu",
        semester="3",
        academic_year="2025-2026",
        status="Active",
    )
    add_student(
        "REG2026002",
        "102",
        "Priya Patel",
        21,
        "ECE",
        email="priya@univ.edu",
        semester="5",
        academic_year="2025-2026",
        status="Active",
    )
    add_student(
        "REG2026003",
        "103",
        "Rohan Verma",
        19,
        "CSE",
        email="rohan@univ.edu",
        semester="3",
        academic_year="2025-2026",
        status="Active",
    )

    create_course(
        course_code="CS301",
        course_name="Algorithms",
        branch="CSE",
        credits=4,
        semester="3",
        academic_year="2025-2026",
        faculty_username="faculty_cse",
        capacity=50,
    )
    create_course(
        course_code="EC501",
        course_name="Digital Signal Processing",
        branch="ECE",
        credits=3,
        semester="5",
        academic_year="2025-2026",
        faculty_username="faculty_ece",
        capacity=40,
    )

    enroll_student("REG2026001", "CS301", semester="3", academic_year="2025-2026")
    enroll_student("REG2026003", "CS301", semester="3", academic_year="2025-2026")
    enroll_student("REG2026002", "EC501", semester="5", academic_year="2025-2026")

    # Attendance records for CS301 and EC501
    record_attendance("REG2026001", "CS301", "2026-10-01", status="Present", session="Lecture")
    record_attendance("REG2026001", "CS301", "2026-10-02", status="Late", session="Lecture")
    record_attendance("REG2026003", "CS301", "2026-10-01", status="Absent", session="Lecture")
    record_attendance("REG2026003", "CS301", "2026-10-02", status="Absent", session="Lecture")
    record_attendance("REG2026002", "EC501", "2026-10-01", status="Present", session="Lecture")

    # Grades for CS301 and EC501
    record_grade("REG2026001", "CS301", assessment_type="Overall", marks_obtained=92, max_marks=100)
    record_grade("REG2026003", "CS301", assessment_type="Overall", marks_obtained=64, max_marks=100)
    record_grade("REG2026002", "EC501", assessment_type="Overall", marks_obtained=85, max_marks=100)

    yield
    set_custom_llm_caller(None)
    os.environ["AI_ASSISTANT_PROVIDER"] = "deterministic"


@pytest.fixture
def client():
    app.config["TESTING"] = True
    with app.test_client() as test_client:
        yield test_client


def _login(client, username, password):
    resp = client.post(
        "/api/v1/auth/login",
        json={"username": username, "password": password},
    )
    assert resp.status_code == 200
    return resp.get_json()["csrf_token"]


def test_advanced_analytics_calculations_against_known_dataset(client):
    _login(client, "admin", "Admin@123")

    resp = client.get("/api/v1/analytics/advanced")
    assert resp.status_code == 200
    data = resp.get_json()

    assert data["totals"]["total_students"] == 3
    assert data["totals"]["total_courses"] == 2
    assert data["totals"]["total_enrollments"] == 3

    # Verify branch enrollment breakdown (CSE=2, ECE=1)
    by_branch = {b["branch"]: b for b in data["enrollment_by_branch"]}
    assert by_branch["CSE"]["student_count"] == 2
    assert by_branch["ECE"]["student_count"] == 1

    # Verify attendance calculation:
    # Total sessions: 2 Present, 1 Late (0.5), 2 Absent -> (2.5 / 5) * 100 = 50.0%
    att = data["attendance_summary"]
    assert att["total_sessions"] == 5
    assert att["present"] == 2
    assert att["late"] == 1
    assert att["absent"] == 2
    assert att["attendance_percentage"] == 50.0

    # Chronological trends: 2026-10-01 (2 Present, 1 Absent -> 66.67%), 2026-10-02 (1 Late, 1 Absent -> 25.0%)
    trends = {t["date"]: t for t in data["attendance_trends"]}
    assert trends["2026-10-01"]["attendance_percentage"] == 66.67
    assert trends["2026-10-02"]["attendance_percentage"] == 25.0

    # Grade distribution: 92% -> O (10), 85% -> A+ (9), 64% -> B+ (7)
    dist = {d["letter_grade"]: d["count"] for d in data["grade_distribution"]}
    assert dist["O"] == 1
    assert dist["A+"] == 1
    assert dist["B+"] == 1

    # GPA summary across 3 students: (10.0 + 7.0 + 9.0) / 3 = 8.67
    assert data["gpa_summary"]["evaluated_students"] == 3
    assert data["gpa_summary"]["average_cgpa"] == 8.67


def test_pdf_excel_and_csv_generation_and_preview(client):
    csrf = _login(client, "admin", "Admin@123")

    # 1. Preview endpoint
    prev = client.get("/api/v1/reports/preview?report_type=student_directory&branch=CSE")
    assert prev.status_code == 200
    prev_json = prev.get_json()
    assert prev_json["row_count"] == 2
    assert prev_json["applied_filters"]["branch"] == "CSE"

    # 2. CSV export across report types
    for rtype in (
        "student_directory",
        "course_enrollment",
        "attendance_register",
        "grade_sheet",
        "academic_transcript",
        "executive_summary",
    ):
        csv_resp = client.get(f"/api/v1/reports/export?report_type={rtype}&format=csv")
        assert csv_resp.status_code == 200
        assert "text/csv" in csv_resp.headers["Content-Type"]
        assert ".csv" in csv_resp.headers["Content-Disposition"]
        text = csv_resp.data.decode("utf-8-sig")
        assert "Report:" in text

    # 3. Excel (.xlsx) export & workbook structure verification
    xlsx_resp = client.post(
        "/api/v1/reports/export",
        json={"report_type": "executive_summary", "format": "xlsx"},
        headers={"X-CSRF-Token": csrf},
    )
    assert xlsx_resp.status_code == 200
    assert (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        in xlsx_resp.headers["Content-Type"]
    )
    wb = load_workbook(io.BytesIO(xlsx_resp.data))
    assert "Report Data" in wb.sheetnames
    assert "Summary & Metadata" in wb.sheetnames
    assert "Semester Breakdown" in wb.sheetnames
    assert "Course Performance" in wb.sheetnames

    # 4. PDF export & binary header verification
    pdf_resp = client.get("/api/v1/reports/export?report_type=academic_transcript&format=pdf")
    assert pdf_resp.status_code == 200
    assert pdf_resp.headers["Content-Type"] == "application/pdf"
    assert pdf_resp.data.startswith(b"%PDF-")


def test_spreadsheet_formula_injection_and_safe_filename_protection(client):
    csrf = _login(client, "admin", "Admin@123")

    # Unit check on sanitize_spreadsheet_cell
    assert sanitize_spreadsheet_cell("=1+1") == "'=1+1"
    assert sanitize_spreadsheet_cell("+cmd|'/C calc'!A0") == "'+cmd|'/C calc'!A0"
    assert sanitize_spreadsheet_cell("-2+3") == "'-2+3"
    assert sanitize_spreadsheet_cell("@SUM(A1:A5)") == "'@SUM(A1:A5)"
    assert sanitize_spreadsheet_cell("  =HYPERLINK('http://evil')") == "'  =HYPERLINK('http://evil')"
    assert sanitize_spreadsheet_cell(92.5) == 92.5

    # Unit check on build_safe_filename against path traversal
    safe_name = build_safe_filename("../../etc/passwd", "pdf", suffix="../secret")
    assert ".." not in safe_name
    assert "/" not in safe_name
    assert "\\" not in safe_name
    assert safe_name.endswith(".pdf")

    # End-to-end check: record attendance with formula-triggering remark and export CSV & Excel
    record_attendance(
        "REG2026001",
        "CS301",
        "2026-10-05",
        status="Present",
        session="Lecture",
        remarks="=CMD|' /C calc'!A0",
    )

    csv_resp = client.get(
        "/api/v1/reports/export?report_type=attendance_register&format=csv&date_from=2026-10-05&date_to=2026-10-05"
    )
    assert csv_resp.status_code == 200
    rows = list(csv.reader(io.StringIO(csv_resp.data.decode("utf-8-sig"))))
    data_row = rows[-1]
    assert data_row[-1] == "'=CMD|' /C calc'!A0"


def test_empty_reports_invalid_filters_and_limits(client):
    _login(client, "admin", "Admin@123")

    # Empty date window returns 200 with 0 rows (no crash)
    empty_prev = client.get(
        "/api/v1/reports/preview?report_type=attendance_register&date_from=2026-01-01&date_to=2026-01-10"
    )
    assert empty_prev.status_code == 200
    assert empty_prev.get_json()["row_count"] == 0

    # Invalid report_type -> 400
    assert (
        client.get("/api/v1/reports/preview?report_type=drop_database").status_code
        == 400
    )

    # Invalid export format -> 400
    assert (
        client.get("/api/v1/reports/export?report_type=student_directory&format=exe").status_code
        == 400
    )

    # Invalid date_from > date_to -> 400
    assert (
        client.get(
            "/api/v1/reports/preview?report_type=attendance_register&date_from=2026-10-10&date_to=2026-10-01"
        ).status_code
        == 400
    )

    # Date range exceeding 366 days -> 400
    assert (
        client.get(
            "/api/v1/reports/preview?report_type=attendance_register&date_from=2024-01-01&date_to=2026-10-01"
        ).status_code
        == 400
    )

    # Limit exceeding MAX_EXPORT_ROWS (2000) -> 400
    assert (
        client.get(
            "/api/v1/reports/preview?report_type=student_directory&limit=99999"
        ).status_code
        == 400
    )


def test_export_and_analytics_rbac_boundaries(client):
    # 1. Unauthenticated requests -> 401
    assert client.get("/api/v1/analytics/advanced").status_code == 401
    assert client.get("/api/v1/reports/preview").status_code == 401
    assert client.get("/api/v1/reports/export").status_code == 401
    assert client.post("/api/v1/ai/query", json={"question": "Hello"}).status_code == 401

    # 2. Faculty (faculty_cse assigned to CSE)
    _login(client, "faculty_cse", "Faculty@123")
    fac_analytics = client.get("/api/v1/analytics/advanced").get_json()
    # Faculty only sees CSE students (2) and CS301 (1 course)
    assert fac_analytics["totals"]["total_students"] == 2
    assert fac_analytics["totals"]["total_courses"] == 1

    # Faculty blocked from exporting unassigned course EC501 or branch ECE
    assert (
        client.get(
            "/api/v1/reports/export?report_type=grade_sheet&course_code=EC501&format=csv"
        ).status_code
        == 403
    )
    assert (
        client.get(
            "/api/v1/reports/export?report_type=student_directory&branch=ECE&format=csv"
        ).status_code
        == 403
    )

    # 3. Student (student_demo linked to REG2026001)
    _login(client, "student_demo", "Student@123")
    stu_prev = client.get("/api/v1/reports/preview?report_type=academic_transcript")
    assert stu_prev.status_code == 200
    stu_rows = stu_prev.get_json()["rows"]
    assert len(stu_rows) == 1
    assert stu_rows[0]["reg_no"] == "REG2026001"

    # Student blocked from requesting another student's transcript or executive_summary
    assert (
        client.get(
            "/api/v1/reports/export?report_type=academic_transcript&reg_no=REG2026002&format=pdf"
        ).status_code
        == 403
    )
    assert (
        client.get(
            "/api/v1/reports/export?report_type=executive_summary&format=pdf"
        ).status_code
        == 403
    )


def test_ai_assistant_grounding_rbac_mutation_guard_and_provider_failures(client):
    # 1. Admin grounded queries
    csrf = _login(client, "admin", "Admin@123")

    status_resp = client.get("/api/v1/ai/status")
    assert status_resp.status_code == 200
    assert status_resp.get_json()["requires_api_key_for_core"] is False

    # Branch enrollment question
    q1 = client.post(
        "/api/v1/ai/query",
        json={"question": "How many students are enrolled in each branch?"},
        headers={"X-CSRF-Token": csrf},
    )
    assert q1.status_code == 200
    q1_data = q1.get_json()
    assert q1_data["intent"] == "branch_enrollment"
    assert "CSE" in q1_data["answer"]
    assert "ECE" in q1_data["answer"]
    assert q1_data["statistics"]["total_students"] == 3

    # Lowest attendance course question (CS301 has 37.5%, EC501 has 100%)
    q2 = client.post(
        "/api/v1/ai/query",
        json={"question": "Which courses have the lowest attendance?"},
        headers={"X-CSRF-Token": csrf},
    ).get_json()
    assert q2["intent"] == "lowest_attendance_courses"
    assert "CS301" in q2["answer"]
    assert q2["statistics"]["lowest_attendance_courses"][0]["course_code"] == "CS301"

    # Empty filter grounding (does not invent stats)
    q_empty = client.post(
        "/api/v1/ai/query",
        json={
            "question": "Summarize course performance",
            "filters": {"branch": "ME"},
        },
        headers={"X-CSRF-Token": csrf},
    ).get_json()
    assert "No matching records" in q_empty["answer"]
    assert q_empty["statistics"]["total_students"] == 0

    # Mutation guard (Requirement 13: requires explicit confirmation before modifying data)
    q_mut = client.post(
        "/api/v1/ai/query",
        json={"question": "Delete all students with attendance below 75%"},
        headers={"X-CSRF-Token": csrf},
    ).get_json()
    assert q_mut["requires_confirmation"] is True
    assert q_mut["action_blocked"] is True

    # 2. AI RBAC boundaries for Faculty and Student
    fac_csrf = _login(client, "faculty_cse", "Faculty@123")
    fac_forbidden = client.post(
        "/api/v1/ai/query",
        json={"question": "Summarize the academic performance of course EC501"},
        headers={"X-CSRF-Token": fac_csrf},
    )
    assert fac_forbidden.status_code == 403

    stu_csrf = _login(client, "student_demo", "Student@123")
    stu_forbidden_reg = client.post(
        "/api/v1/ai/query",
        json={"question": "What is the GPA for REG2026002?"},
        headers={"X-CSRF-Token": stu_csrf},
    )
    assert stu_forbidden_reg.status_code == 403

    stu_forbidden_course = client.post(
        "/api/v1/ai/query",
        json={"question": "Summarize performance of EC501"},
        headers={"X-CSRF-Token": stu_csrf},
    )
    assert stu_forbidden_course.status_code == 403

    stu_own = client.post(
        "/api/v1/ai/query",
        json={"question": "Summarize my academic performance"},
        headers={"X-CSRF-Token": stu_csrf},
    )
    assert stu_own.status_code == 200
    assert stu_own.get_json()["filters_used"]["reg_no"] == "REG2026001"

    # 3. External AI provider privacy & failure handling
    admin_csrf = _login(client, "admin", "Admin@123")
    os.environ["AI_ASSISTANT_PROVIDER"] = "gemini"
    os.environ.pop("AI_ASSISTANT_API_KEY", None)

    # Missing API key when external provider is enabled -> graceful 503
    missing_key_resp = client.post(
        "/api/v1/ai/query",
        json={"question": "How many students are in each branch?"},
        headers={"X-CSRF-Token": admin_csrf},
    )
    assert missing_key_resp.status_code == 503

    # Verify PII is stripped when external LLM caller is invoked
    captured_contexts = []

    def mock_llm_success(provider, question, context, draft):
        captured_contexts.append(context)
        return f"Verified AI Summary: {draft}"

    set_custom_llm_caller(mock_llm_success)
    ext_ok = client.post(
        "/api/v1/ai/query",
        json={"question": "How many students are in each branch?"},
        headers={"X-CSRF-Token": admin_csrf},
    )
    assert ext_ok.status_code == 200
    assert "Verified AI Summary:" in ext_ok.get_json()["answer"]
    serialized_ctx = str(captured_contexts[0])
    assert "Aarav Sharma" not in serialized_ctx
    assert "REG2026001" not in serialized_ctx

    # Simulate 429 rate-limit from external provider
    def mock_llm_rate_limited(**kwargs):
        raise urllib.error.HTTPError(
            url="https://api.example.com",
            code=429,
            msg="Too Many Requests",
            hdrs={},
            fp=None,
        )

    set_custom_llm_caller(mock_llm_rate_limited)
    rl_resp = client.post(
        "/api/v1/ai/query",
        json={"question": "How many students are in each branch?"},
        headers={"X-CSRF-Token": admin_csrf},
    )
    assert rl_resp.status_code == 429
